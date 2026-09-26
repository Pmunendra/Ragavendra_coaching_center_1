"""Bulk question import (CSV / XLSX).

Import is a two-step preview-then-confirm flow: parse_upload() +
validate_rows() run once to show the admin exactly what will and won't be
imported (with a reason for every failure), and only rows that passed
validation are ever written to the database. Nothing is imported silently
or partially - a row is either clean or it's skipped with a visible reason.
"""
import csv
import io
import uuid
from dataclasses import dataclass, field

from openpyxl import load_workbook

from app.models import QUESTION_TYPES

TEMPLATE_COLUMNS = [
    "question_text", "question_type", "option_a", "option_b", "option_c", "option_d",
    "correct_option", "marks", "negative_marks", "subject", "explanation",
]

VALID_QUESTION_TYPES = {t[0] if isinstance(t, tuple) else t for t in QUESTION_TYPES}


@dataclass
class ImportRow:
    row_number: int
    raw: dict
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)

    @property
    def is_valid(self):
        return not self.errors


def generate_template_csv():
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(TEMPLATE_COLUMNS)
    writer.writerow([
        "Capital of India?", "single_choice", "Mumbai", "New Delhi", "Kolkata", "Chennai",
        "B", "1", "0.25", "General Knowledge", "New Delhi is the capital of India.",
    ])
    writer.writerow([
        "Select all prime numbers", "multiple_choice", "2", "3", "4", "9",
        "A,B", "2", "0", "Mathematics", "",
    ])
    return buf.getvalue()


def parse_upload(file_storage):
    """Returns list[dict] of raw rows, or raises ValueError with a
    human-readable message if the file can't be read at all."""
    filename = (file_storage.filename or "").lower()
    data = file_storage.read()

    if filename.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        rows = [{(k or "").strip().lower(): v for k, v in r.items()} for r in reader]
    elif filename.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(data), data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = [str(h).strip().lower() if h is not None else "" for h in next(rows_iter)]
        except StopIteration:
            return []
        rows = []
        for r in rows_iter:
            if all(c is None or str(c).strip() == "" for c in r):
                continue
            rows.append({header[i]: (r[i] if i < len(r) else None) for i in range(len(header))})
    else:
        raise ValueError("Unsupported file type - please upload a .csv or .xlsx file.")

    return rows


def _clean(value):
    if value is None:
        return ""
    return str(value).strip()


def validate_rows(raw_rows, exam):
    """Validates every row against the exam it will be imported into.
    Checks: required fields present, at least 2 non-empty options, a
    correct_option that actually references one of the given options,
    valid question_type, numeric marks, duplicate question text (both
    within the file and already existing in this exam)."""
    from app.models import Subject, Question

    existing_texts = {
        (q.question_text or "").strip().lower()
        for q in exam.questions
    } if hasattr(exam, "questions") else set()

    subjects_by_name = {s.name.strip().lower(): s for s in Subject.query.all()}

    seen_in_file = set()
    results = []

    for i, raw in enumerate(raw_rows, start=2):  # row 1 is the header
        row = ImportRow(row_number=i, raw=raw)
        norm = {k.strip().lower(): _clean(v) for k, v in raw.items()}

        question_text = norm.get("question_text", "")
        if not question_text:
            row.errors.append("Question text is required.")

        q_type = (norm.get("question_type") or "single_choice").strip().lower()
        if q_type not in VALID_QUESTION_TYPES:
            row.errors.append(f"Unknown question_type '{q_type}'.")

        options = []
        for letter in ("a", "b", "c", "d", "e", "f"):
            val = norm.get(f"option_{letter}", "")
            if val:
                options.append((letter.upper(), val))
        if q_type in ("single_choice", "multiple_choice") and len(options) < 2:
            row.errors.append("At least 2 options are required.")

        correct_raw = norm.get("correct_option", "")
        if q_type in ("single_choice", "multiple_choice"):
            if not correct_raw:
                row.errors.append("correct_option is required (e.g. 'B' or 'A,C').")
            else:
                correct_letters = {c.strip().upper() for c in correct_raw.split(",") if c.strip()}
                option_letters = {letter for letter, _ in options}
                unknown = correct_letters - option_letters
                if unknown:
                    row.errors.append(f"correct_option references option(s) not provided: {', '.join(sorted(unknown))}.")
                if q_type == "single_choice" and len(correct_letters) > 1:
                    row.errors.append("single_choice questions must have exactly one correct_option.")

        marks_raw = norm.get("marks", "1") or "1"
        try:
            marks = float(marks_raw)
            if marks <= 0:
                row.errors.append("marks must be a positive number.")
        except ValueError:
            row.errors.append(f"marks '{marks_raw}' is not a valid number.")
            marks = 1

        neg_raw = norm.get("negative_marks", "0") or "0"
        try:
            negative_marks = float(neg_raw)
        except ValueError:
            row.errors.append(f"negative_marks '{neg_raw}' is not a valid number.")
            negative_marks = 0

        subject_name = norm.get("subject", "")
        subject_obj = subjects_by_name.get(subject_name.lower()) if subject_name else None
        if subject_name and not subject_obj:
            row.warnings.append(f"Subject '{subject_name}' doesn't exist yet - will be imported without a subject.")

        norm_text = question_text.strip().lower()
        if norm_text:
            if norm_text in existing_texts:
                row.errors.append("This question already exists in this exam (duplicate).")
            elif norm_text in seen_in_file:
                row.errors.append("Duplicate of another question earlier in this file.")
            else:
                seen_in_file.add(norm_text)

        row.parsed = dict(
            question_text=question_text, question_type=q_type, options=options,
            correct_letters=(correct_raw and {c.strip().upper() for c in correct_raw.split(",") if c.strip()}) or set(),
            marks=marks, negative_marks=negative_marks,
            subject_id=subject_obj.id if subject_obj else None,
            explanation=norm.get("explanation") or None,
        )
        results.append(row)

    return results


def commit_valid_rows(validated_rows, exam):
    """Inserts only the rows that passed validation. Returns the count
    actually imported."""
    from app.extensions import db
    from app.models import Question, Option

    start_index = exam.question_count()
    imported = 0
    for row in validated_rows:
        if not row.is_valid:
            continue
        p = row.parsed
        q = Question(
            exam_id=exam.id,
            subject_id=p["subject_id"],
            question_type=p["question_type"],
            question_text=p["question_text"],
            marks=p["marks"],
            negative_marks=p["negative_marks"],
            explanation=p["explanation"],
            order_index=start_index + imported,
        )
        db.session.add(q)
        db.session.flush()
        for letter, text in p["options"]:
            db.session.add(Option(
                question_id=q.id, option_text=text,
                is_correct=letter in p["correct_letters"],
                order_index=ord(letter) - ord("A"),
            ))
        imported += 1

    if imported:
        exam.recompute_total_marks()
    return imported
